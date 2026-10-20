"""Реестр ЮKassa (A5), сверка выплат с выпиской и КУДиР (A7). ЮKassa — подменённая, только GET."""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from sqlalchemy import text

from app.core import heartbeat
from app.core.config import settings
from app.messages import manager as manager_messages
from app.modules.payment import kudir, registry, yookassa_client

NOW = datetime(2026, 11, 20, 6, 30, tzinfo=timezone.utc)  # 09:30 МСК, пятница


def payment(pid, amount, income, captured, method="bank_card", status="succeeded", test=False):
    return {"id": pid, "status": status, "amount": {"value": f"{amount:.2f}", "currency": "RUB"},
            "income_amount": {"value": f"{income:.2f}", "currency": "RUB"}, "test": test,
            "payment_method": {"type": method}, "created_at": captured, "captured_at": captured}


PAYMENTS = [
    payment("p1", 1000, 965, "2026-11-16T09:00:00.000Z"),             # пн
    payment("p2", 2000, 1930, "2026-11-16T15:00:00.000Z", "sbp"),     # пн
    payment("p3", 500, 482.5, "2026-11-17T10:00:00.000Z"),            # вт — выплаты не будет
    payment("p4", 700, 0, "2026-11-17T10:00:00.000Z", status="canceled"),
    payment("t1", 900, 870, "2026-11-16T10:00:00.000Z", test=True),   # тестовый магазин — мимо реестра
]
REFUNDS = [{"id": "r1", "payment_id": "p2", "status": "succeeded", "amount": {"value": "300.00"},
            "created_at": "2026-11-16T18:00:00.000Z"}]


@pytest.fixture
def yookassa(monkeypatch):
    calls = []

    async def fake_call(method, path, payload=None, key=""):
        calls.append((method, path))
        assert method == "GET"
        return {"items": PAYMENTS if path.startswith("/payments") else REFUNDS, "next_cursor": None}

    monkeypatch.setattr(yookassa_client, "_call", fake_call)
    monkeypatch.setattr(settings, "yookassa_shop_id", "1")
    monkeypatch.setattr(settings, "yookassa_secret_key", "k")
    return calls


@pytest.fixture
def alerts(monkeypatch):
    sent = []

    async def notify(kind, text_, **kwargs):
        sent.append(text_)
        return True

    monkeypatch.setattr(manager_messages, "notify", notify)
    return sent


@pytest.fixture
async def bank(clean):
    """Выписка: в тестовой базе cashflow нет — кладём две таблицы, после теста убираем."""
    async with clean() as session:
        await session.execute(text("delete from client_notices"))
        created = not (await session.execute(text("select to_regclass('finance.fact_operations') is not null"))).scalar()
        if created:
            await session.execute(text(
                "create table finance.articles (article_id text primary key, article text, grp text, section text,"
                " is_capex boolean default false)"))
            await session.execute(text(
                "create table finance.fact_operations (operation_id text, split_no smallint default 1,"
                " operation_date date, direction text, amount numeric, article_id text)"))
            await session.execute(text(
                "insert into finance.articles values ('cfo_in_yookassa', 'Поступления ЮKassa', 'Выручка', 'CFO'),"
                " ('cfo_in_direct', 'Прямые оплаты покупателей', 'Выручка', 'CFO'),"
                " ('cff_owner_contribution', 'Вклад собственника', 'Собственники', 'CFF')"))
        for table in ("yk_payments", "yk_refunds", "yk_payouts"):
            await session.execute(text(f"delete from finance.{table}"))
        await session.commit()
        await kudir.ensure_view(await session.connection())
        await session.commit()
    yield clean
    async with clean() as session:
        if created:
            await session.execute(text("drop view if exists finance.v_yk_kudir"))
            await session.execute(text("drop table finance.fact_operations, finance.articles"))
            await kudir.ensure_view(await session.connection())
        await session.commit()


async def bank_rows(db, *rows):
    async with db() as session:
        for op, day, amount, article in rows:
            await session.execute(text(
                "insert into finance.fact_operations values (:o, 1, :d, 'in', :a, :art)"),
                {"o": op, "d": day, "a": amount, "art": article})
        await session.commit()


async def select(db, sql):
    async with db() as session:
        return [dict(r._mapping) for r in await session.execute(text(sql))]


async def test_registry_stores_payments_with_fee_and_skips_test(bank, yookassa, alerts):
    await bank_rows(bank, ("b1", date(2026, 11, 17), Decimal("2595.00"), "cfo_in_yookassa"))
    result = await registry.sync(force=True, now=NOW)
    assert result["платежей"] == 4 and result["возвратов"] == 1
    assert all(method == "GET" for method, _ in yookassa)
    rows = {r["id"]: r for r in await select(bank, "select * from finance.yk_payments")}
    assert "t1" not in rows
    assert rows["p2"]["fee"] == Decimal("70.00") and rows["p2"]["payment_method"] == "sbp"
    # Выплата 17.11 = 965 + 1930 − 300 за понедельник.
    [payout] = await select(bank, "select * from finance.yk_payouts")
    assert payout["status"] == "matched" and payout["days_from"] == payout["days_to"] == date(2026, 11, 16)
    paid = await select(bank, "select id, payout_date from finance.yk_payments order by id")
    assert {r["id"]: r["payout_date"] for r in paid} == {
        "p1": date(2026, 11, 17), "p2": date(2026, 11, 17), "p3": None, "p4": None}
    assert result["в пути, ₽"] == "482.50"
    # 17.11 → 20.11 — три рабочих дня, ещё не тревога.
    assert result["не выплачено дольше 3 рабочих дней"] == 0 and alerts == []
    assert await heartbeat.last_run(registry.HEARTBEAT) is not None


async def test_mismatch_and_overdue_are_reported_once(bank, yookassa, alerts):
    await bank_rows(bank, ("b1", date(2026, 11, 17), Decimal("2500.00"), "cfo_in_yookassa"))
    later = datetime(2026, 11, 24, 7, 0, tzinfo=timezone.utc)
    result = await registry.sync(force=True, now=later)
    assert result["расхождений"] == 1 and result["не выплачено дольше 3 рабочих дней"] == 2
    assert len(alerts) == 3
    assert "не сходится" in alerts[0] and "разница -95.00 ₽" in alerts[0]
    await registry.sync(force=True, now=later)
    assert len(alerts) == 3


async def test_payout_covering_a_weekend(bank, yookassa, alerts):
    # Выплата понедельника 23.11 покрывает и пн 16.11, и вт 17.11 — дни подряд.
    await bank_rows(bank, ("b1", date(2026, 11, 23), Decimal("3077.50"), "cfo_in_yookassa"))
    result = await registry.sync(force=True, now=datetime(2026, 11, 23, 9, 0, tzinfo=timezone.utc))
    assert result["сошлось"] == 1 and result["в пути, ₽"] == "0"
    [payout] = await select(bank, "select days_from, days_to from finance.yk_payouts")
    assert payout == {"days_from": date(2026, 11, 16), "days_to": date(2026, 11, 17)}


async def test_daily_schedule(bank, yookassa, alerts, monkeypatch):
    early = datetime(2026, 11, 20, 3, 0, tzinfo=timezone.utc)  # 06:00 МСК
    assert (await registry.sync(now=early))["skipped"] == "рано: ждём выписку"
    assert "платежей" in await registry.sync(now=NOW)

    async def ran_today(name):
        return NOW

    monkeypatch.setattr(heartbeat, "last_run", ran_today)
    assert (await registry.sync(now=NOW))["skipped"] == "сегодня уже"


async def test_kudir_rows_refunds_and_bank_income(bank, yookassa, alerts):
    await bank_rows(bank, ("b1", date(2026, 11, 17), Decimal("2595.00"), "cfo_in_yookassa"),
                    ("b2", date(2026, 11, 18), Decimal("400.00"), "cfo_in_direct"),
                    ("b3", date(2026, 11, 18), Decimal("50000.00"), "cff_owner_contribution"))
    await registry.sync(force=True, now=NOW)
    records = await kudir.rows(date(2026, 11, 1), date(2026, 11, 30))
    assert [(r["doc_number"], r["income"]) for r in records] == [
        ("p1", Decimal("1000.00")), ("p2", Decimal("2000.00")), ("r1", Decimal("-300.00")),
        ("p3", Decimal("500.00")), ("b2", Decimal("400.00"))]
    # Выплата ЮKassa (нетто) и вклад собственника доходом не считаются.
    assert sum(r["income"] for r in records) == Decimal("3600.00")
    assert "через ЮKassa (sbp)" in records[1]["content"]
    csv_text = kudir.to_csv(records).decode("utf-8-sig")
    assert "Итого;;;3600,00" in csv_text
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(kudir.to_xlsx(records, date(2026, 11, 1), date(2026, 11, 30))))
    sheet = book.active
    assert "подтвердить у бухгалтера" in sheet["A2"].value
    assert sheet.cell(row=sheet.max_row, column=5).value == 3600.0


async def test_kudir_by_payout_date(bank, yookassa, alerts, monkeypatch):
    monkeypatch.setattr(settings, "kudir_income_date", "payout_date")
    await bank_rows(bank, ("b1", date(2026, 11, 17), Decimal("2595.00"), "cfo_in_yookassa"))
    await registry.sync(force=True, now=NOW)
    async with bank() as session:
        await kudir.ensure_view(await session.connection())
        await session.commit()
    records = await kudir.rows(date(2026, 11, 1), date(2026, 11, 30))
    # Невыплаченный p3 дохода ещё не дал; выплаченные — датой выплаты.
    assert {r["doc_number"]: r["income_date"] for r in records if r["kind"] != "bank"} == {
        "p1": date(2026, 11, 17), "p2": date(2026, 11, 17), "r1": date(2026, 11, 17)}
    monkeypatch.setattr(settings, "kudir_income_date", "captured_at")
    async with bank() as session:
        await kudir.ensure_view(await session.connection())
        await session.commit()


def test_kudir_period_parsing():
    today = date(2026, 11, 3)
    assert kudir.parse_period([], today) == (date(2026, 10, 1), date(2026, 10, 31))
    assert kudir.parse_period(["2026-12"], today) == (date(2026, 12, 1), date(2026, 12, 31))
    assert kudir.parse_period(["2027"], today) == (date(2027, 1, 1), date(2027, 12, 31))
    assert kudir.parse_period(["2026-10-01", "2026-12-31"], today) == (date(2026, 10, 1), date(2026, 12, 31))
    with pytest.raises(ValueError):
        kudir.parse_period(["вчера"], today)


async def test_kudir_command_sends_two_files(bank, yookassa, alerts, monkeypatch):
    from app.modules.dialog import telegram_client
    from app.modules.ops import commands

    monkeypatch.setattr(settings, "ops_admin_tg_ids", "")
    monkeypatch.setattr(settings, "ops_assistant_allowed_tg_ids", "")
    sent, files = [], []

    async def reply(text_, **kwargs):
        sent.append(text_)

    async def send_document(filename, content, **kwargs):
        files.append(filename)

    monkeypatch.setattr(telegram_client, "reply", reply)
    monkeypatch.setattr(telegram_client, "send_document", send_document)
    await registry.sync(force=True, now=NOW)
    message = {"text": "/kudir 2026-11", "chat": {"id": -1}, "message_id": 3, "from": {"id": 5}}
    assert await commands.handle(message, {"username": "bot"})
    assert "записей 4" in sent[0] and "бухгалтер" in sent[0]
    assert files == ["kudir_20261101_20261130.xlsx", "kudir_20261101_20261130.csv"]


async def test_refund_bigger_than_income_is_not_an_unpaid_day(bank, monkeypatch, alerts):
    """26.09.2026: оплата 217 ₽ и полный возврат в тот же день — нетто −8,25 ₽, это не невыплата."""
    async def fake_call(method, path, payload=None, key=""):
        if path.startswith("/payments"):
            return {"items": [payment("s1", 217, 208.75, "2026-09-26T10:00:00.000Z", "sberbank")]}
        return {"items": [{"id": "r9", "payment_id": "s1", "status": "succeeded", "amount": {"value": "217.00"},
                           "created_at": "2026-09-26T12:00:00.000Z"}]}

    monkeypatch.setattr(yookassa_client, "_call", fake_call)
    monkeypatch.setattr(settings, "yookassa_shop_id", "1")
    monkeypatch.setattr(settings, "yookassa_secret_key", "k")
    result = await registry.sync(force=True, now=datetime(2026, 10, 20, 19, 0, tzinfo=timezone.utc))
    assert result["не выплачено дольше 3 рабочих дней"] == 0 and alerts == []
    assert result["в пути, ₽"] == "-8.25"
